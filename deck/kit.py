"""Slide kit: reusable slide types in the workshop deck's style.

A deck is a list of dicts: {"kind": "cards", "title": ..., ...}. `build_deck()` opens the real deck
(for its layouts, fonts and footer), clears the slides, and renders each dict with the matching
function in TEMPLATES. Content lives in decks/*.py; this file only knows layout.
"""
from __future__ import annotations

import copy
import json
import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from build import (  # noqa: F401  (shared helpers and palette)
    AMBER, BODY, CYAN, HEAD, INK, MUTED, NAVY, NAVY_2, PANEL, PURPLE, RULE, SLATE, TINT, WHITE,
    bullets, chip, delete_slide, notes, rgb, set_title, shape, text,
)

HERE = Path(__file__).parent
SRC = HERE / "source" / "workshop-deck.pptx"
ICON_DIR = HERE / "icons" / "all"
GREEN, RED, SOFT = "10B981", "EF4444", "CBD5E1"
MONO = "Courier New"

ICON = {  # name -> file stem in icons/all (colour is baked into each icon)
    "phone_w": "0cfb684d", "search": "12be325d", "plane": "164d0579", "cap": "17e4a0ed", "hand_w": "28a6d33f",
    "dollar": "29c6a7d9", "bolt": "366e1990", "palette": "36ebcd7a", "rocket": "3bdf0d77", "chat": "4036ec40",
    "check": "4f1c2c6d", "hand": "5210d3a4", "code": "534b8949", "bulb": "54383f80", "globe": "60298cbe",
    "people": "65bca282", "mail_w": "6e8c60bd", "x": "7ef058e5", "doc": "94ed0060", "in_w": "b7dfbe58",
    "globe_w": "b8107fb1", "note": "be33ba18", "in": "c136a8a2", "chart": "cf5e672d", "cloud": "d256c8e7",
    "target": "d8fc09b6", "layers": "d9e5fb61", "rocket_w": "df142025", "codefile": "e32f45bf", "mobile": "e467b63b",
}
_NUMBER_SP = None  # slide-number placeholder copied from the source deck


# ----------------------------------------------------------------------------- primitives
def icon_file(name: str) -> Path:
    stem = ICON[name]
    for p in ICON_DIR.glob(stem + ".*"):
        return p
    raise FileNotFoundError(name)


def icon_in_circle(s, x, y, d, name, circle):
    shape(s, MSO_SHAPE.OVAL, x, y, d, d, circle, name=f"{name} circle")
    pad = d * 0.27
    pic = s.shapes.add_picture(str(icon_file(name)), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))
    pic.name = f"icon:{name}"


def alpha(sh, pct):
    clr = sh._element.spPr.find(qn("a:solidFill"))[0]
    etree.SubElement(clr, qn("a:alpha"), val=str(int(pct * 1000)))


def add_number(s):
    if _NUMBER_SP is not None:
        s.shapes._spTree.append(copy.deepcopy(_NUMBER_SP))


def new_slide(prs, layout):
    return prs.slides.add_slide({l.name: l for l in prs.slide_layouts}[layout])


def light(prs, spec):
    """Content frame: title, optional subtitle, slide number, notes."""
    s = new_slide(prs, "ContentLight")
    set_title(s, spec["title"])
    if spec.get("subtitle"):
        text(s, 0.7, 1.3, 12.0, 0.4, dict(t=spec["subtitle"], size=16, color=SLATE), name="Subtitle")
    add_number(s)
    if spec.get("notes"):
        notes(s, spec["notes"])
    return s


def dark(prs, spec, *, circles=True):
    s = new_slide(prs, "TitleDark")
    if circles:
        shape(s, MSO_SHAPE.OVAL, 10.9, -2.5, 5.2, 5.2, PURPLE, name="Decor purple")
        c = shape(s, MSO_SHAPE.OVAL, -2.6, 5.6, 5.0, 5.0, CYAN, name="Decor cyan")
        alpha(c, 22)
    if spec.get("notes"):
        notes(s, spec["notes"])
    return s


def source_line(s, src):
    if src:
        text(s, 0.7, 6.88, 11.5, 0.25, dict(t="Source: " + src, size=10, color=MUTED), name="Source")


# ----------------------------------------------------------------------------- slide types
def cover(prs, sp):
    s = dark(prs, sp)
    text(s, 0.9, 1.75, 9.0, 0.4, dict(t=sp["eyebrow"].upper(), size=14, bold=True, color=CYAN), name="Eyebrow")
    set_title(s, sp["title"], size=46, color=WHITE, x=0.8, y=2.35, w=11.6, h=1.7)
    text(s, 0.9, 4.2, 10.5, 1.0, dict(t=sp["subtitle"], size=20, color=SOFT), name="Subtitle")
    icon_in_circle(s, 0.9, 5.55, 0.62, "rocket_w", "162238")
    text(s, 1.7, 5.55, 9.0, 0.65, [dict(t="Adamya Kumar", size=14, bold=True, color=WHITE), dict(t=sp["part"], size=13, color=MUTED)], name="Presenter")
    return s


def section(prs, sp):
    s = dark(prs, sp)
    text(s, 0.9, 1.4, 4.0, 2.6, dict(t=sp["number"], size=150, bold=True, font=HEAD, color=NAVY), name="Number")
    set_title(s, sp["title"], size=40, color=WHITE, x=0.8, y=3.9, w=11.5, h=1.0)
    text(s, 0.9, 4.95, 10.5, 1.2, dict(t=sp["subtitle"], size=20, color=SOFT), name="Subtitle")
    return s


def cards(prs, sp):
    s = light(prs, sp)
    items = sp["items"]
    n = len(items)
    cols = sp.get("cols") or (n if n <= 4 else 3)
    rows = math.ceil(n / cols)
    top = 1.85 if sp.get("subtitle") else 1.7
    gap = 0.25
    w = (12.15 - (cols - 1) * gap) / cols
    h = (6.75 - top - (rows - 1) * gap) / rows
    if rows == 1:
        h = min(h, 4.3)
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x, y = 0.6 + c * (w + gap), top + r * (h + gap)
        pattern = {3: [0, 1, 0], 4: [0, 1, 1, 0]}.get(n)
        d = (pattern[i] if pattern else (r + c) % 2) == 1
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, NAVY if d else TINT, radius=0.07, shadow=True, name=f"Card {i + 1}")
        ink, sub = (WHITE, SOFT) if d else (INK, SLATE)
        if rows == 1:
            icon_in_circle(s, x + 0.3, y + 0.3, 0.7, it["icon"], NAVY_2 if d else WHITE)
            text(s, x + 0.3, y + 1.2, w - 0.6, 0.9, dict(t=it["head"], size=19, bold=True, font=HEAD, color=ink), name=f"Card {i + 1} head")
            text(s, x + 0.3, y + 2.05, w - 0.6, h - 2.2, dict(t=it["body"], size=14, color=sub), name=f"Card {i + 1} body")
        else:
            icon_in_circle(s, x + 0.25, y + 0.25, 0.6, it["icon"], NAVY_2 if d else WHITE)
            text(s, x + 1.05, y + 0.28, w - 1.25, 0.6, dict(t=it["head"], size=17, bold=True, font=HEAD, color=ink), name=f"Card {i + 1} head", anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.3, y + 1.05, w - 0.6, h - 1.15, dict(t=it["body"], size=14, color=sub), name=f"Card {i + 1} body")
    return s


def compare(prs, sp):
    s = light(prs, sp)
    rowsd = sp["rows"]
    n = len(rowsd)
    top = 1.85
    rh = min(1.05, (6.75 - top - 0.95) / n)
    H = 0.95 + n * rh
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, top, 5.85, H, TINT, radius=0.05, shadow=True, name="Weak panel")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.9, top, 5.85, H, NAVY, radius=0.05, shadow=True, name="Strong panel")
    chip(s, 0.9, top + 0.25, 1.5, sp.get("left_head", "WEAK"), fill=RED, color=WHITE, h=0.36, size=12, name="Weak chip")
    chip(s, 7.2, top + 0.25, 1.5, sp.get("right_head", "STRONG"), fill=GREEN, color=WHITE, h=0.36, size=12, name="Strong chip")
    tri = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 6.5, top + H / 2 - 0.15, 0.3, 0.3, AMBER, name="Arrow")
    tri.rotation = 90
    for i, (a, b) in enumerate(rowsd):
        y = top + 0.8 + i * rh
        text(s, 0.9, y, 5.3, rh - 0.1, dict(t=a, size=15, color=SLATE), name=f"Weak {i + 1}")
        text(s, 7.2, y, 5.3, rh - 0.1, dict(t=b, size=15, bold=True, color=WHITE), name=f"Strong {i + 1}")
    return s


def steps(prs, sp):
    """Horizontal path, 3 to 5 nodes. Items: head, body, optional icon (else a number)."""
    s = light(prs, sp)
    items = sp["items"]
    n = len(items)
    colw = 12.15 / n
    D, Y = 1.05, 2.1
    first, last = 0.6 + D / 2, 0.6 + (n - 1) * colw + D / 2
    shape(s, MSO_SHAPE.RECTANGLE, first, Y + D / 2 - 0.03, last - first, 0.06, SOFT, name="Path line")
    for i, it in enumerate(items):
        x = 0.6 + i * colw
        if i < n - 1:
            t = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + colw / 2 + D / 4, Y + D / 2 - 0.13, 0.26, 0.26, MUTED, name=f"Arrow {i + 1}")
            t.rotation = 90
        if it.get("icon"):
            icon_in_circle(s, x, Y, D, it["icon"], NAVY)
            chip(s, x + D - 0.28, Y - 0.1, 0.4, str(i + 1), h=0.4, size=13, name=f"Step {i + 1}")
        else:
            c = shape(s, MSO_SHAPE.OVAL, x, Y, D, D, NAVY, name=f"Step {i + 1}")
            tf = c.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(i + 1)
            r.font.name, r.font.size, r.font.bold = HEAD, Pt(28), True
            r.font.color.rgb = rgb(WHITE)
        text(s, x, Y + D + 0.3, colw - 0.3, 0.7, dict(t=it["head"], size=20, bold=True, font=HEAD), name=f"Step {i + 1} head")
        text(s, x, Y + D + 1.05, colw - 0.3, 1.9, dict(t=it["body"], size=15, color=SLATE), name=f"Step {i + 1} body")
    if sp.get("callout"):
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 6.0, 12.15, 0.75, TINT, radius=0.18, shadow=True, name="Callout band")
        chip(s, 0.9, 6.2, 1.35, sp.get("callout_tag", "TIP"), h=0.35, size=12, name="Callout tag")
        text(s, 2.5, 6.05, 10.0, 0.65, dict(t=sp["callout"], size=16), name="Callout", anchor=MSO_ANCHOR.MIDDLE)
    return s


def vsteps(prs, sp):
    """Vertical rows, 3 to 6. Items: head, body."""
    s = light(prs, sp)
    items = sp["items"]
    n = len(items)
    top, gap = 1.85, 0.12
    h = (6.75 - top - (n - 1) * gap) / n
    for i, it in enumerate(items):
        y = top + i * (h + gap)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, y, 12.15, h, TINT, radius=0.12, shadow=True, name=f"Row {i + 1}")
        c = shape(s, MSO_SHAPE.OVAL, 0.85, y + h / 2 - 0.25, 0.5, 0.5, NAVY, name=f"Row {i + 1} number")
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        r.font.name, r.font.size, r.font.bold = HEAD, Pt(17), True
        r.font.color.rgb = rgb(WHITE)
        text(s, 1.65, y, 3.3, h, dict(t=it["head"], size=18, bold=True, font=HEAD), name=f"Row {i + 1} head", anchor=MSO_ANCHOR.MIDDLE)
        text(s, 5.1, y, 7.4, h, dict(t=it["body"], size=15, color=SLATE), name=f"Row {i + 1} body", anchor=MSO_ANCHOR.MIDDLE)
    return s


def bullets_stat(prs, sp):
    s = light(prs, sp)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.85, 7.5, 4.85, TINT, radius=0.05, shadow=True, name="Bullets card")
    bullets(s, 0.95, 2.15, 6.8, 4.3, sp["bullets"], size=17, gap=9, name="Bullets")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.35, 1.85, 4.4, 4.85, NAVY, radius=0.05, shadow=True, name="Stat card")
    text(s, 8.6, 2.3, 3.9, 1.6, dict(t=sp["stat"], size=72, bold=True, font=HEAD, color=AMBER, align=PP_ALIGN.CENTER), name="Stat")
    text(s, 8.6, 4.0, 3.9, 0.5, dict(t=sp["stat_label"], size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER), name="Stat label")
    text(s, 8.7, 4.65, 3.7, 1.8, dict(t=sp.get("stat_note", ""), size=13, color=SOFT, align=PP_ALIGN.CENTER), name="Stat note")
    source_line(s, sp.get("source"))
    return s


def two_col(prs, sp):
    s = light(prs, sp)
    for i, col in enumerate((sp["left"], sp["right"])):
        x = 0.6 + i * 6.2
        d = i == 1
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.85, 5.95, 4.85, NAVY if d else TINT, radius=0.05, shadow=True, name=f"Column {i + 1}")
        icon_in_circle(s, x + 0.3, 2.1, 0.65, col["icon"], NAVY_2 if d else WHITE)
        text(s, x + 1.15, 2.1, 4.6, 0.65, dict(t=col["head"], size=20, bold=True, font=HEAD, color=WHITE if d else INK), name=f"Column {i + 1} head", anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, x + 0.35, 3.0, 5.3, 3.6, col["bullets"], size=15, color=WHITE if d else INK, gap=7, name=f"Column {i + 1} bullets")
    return s


def code(prs, sp):
    s = light(prs, sp)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.85, 7.3, 4.85, INK, radius=0.04, shadow=True, name="Code panel")
    chip(s, 0.9, 2.05, 1.3, sp.get("lang", "PYTHON"), fill=PURPLE, color=WHITE, h=0.3, size=11, name="Lang")
    lines = [dict(t=l if isinstance(l, str) else l[0], size=15, font=MONO, color=(SOFT if isinstance(l, str) else l[1])) for l in sp["code"]]
    text(s, 0.95, 2.55, 6.8, 4.0, lines, name="Code")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.15, 1.85, 4.6, 4.85, TINT, radius=0.05, shadow=True, name="Notes card")
    text(s, 8.45, 2.05, 4.0, 0.5, dict(t=sp["side_head"], size=19, bold=True, font=HEAD), name="Side head")
    bullets(s, 8.45, 2.7, 4.05, 3.9, sp["side"], size=15, gap=8, name="Side bullets")
    return s


def table(prs, sp):
    s = light(prs, sp)
    head, rows, widths = sp["head"], sp["rows"], sp["widths"]
    top = 1.85
    n = len(rows)
    rh = min(0.85, (6.75 - top - 0.55) / n)
    x0 = 0.6
    shape(s, MSO_SHAPE.RECTANGLE, x0, top, sum(widths), 0.55, NAVY, name="Header row")
    x = x0
    for j, h in enumerate(head):
        text(s, x + 0.15, top, widths[j] - 0.3, 0.55, dict(t=h, size=14, bold=True, color=WHITE), name=f"Head {j + 1}", anchor=MSO_ANCHOR.MIDDLE)
        x += widths[j]
    for i, row in enumerate(rows):
        y = top + 0.55 + i * rh
        shape(s, MSO_SHAPE.RECTANGLE, x0, y, sum(widths), rh, TINT if i % 2 == 0 else WHITE, name=f"Row {i + 1}")
        x = x0
        for j, cell in enumerate(row):
            text(s, x + 0.15, y, widths[j] - 0.3, rh, dict(t=cell, size=14, bold=(j == 0), color=INK if j == 0 else SLATE), name=f"Cell {i + 1}.{j + 1}", anchor=MSO_ANCHOR.MIDDLE)
            x += widths[j]
    source_line(s, sp.get("source"))
    return s


def activity(prs, sp):
    s = dark(prs, sp, circles=False)
    chip(s, 0.9, 0.9, 2.6, f"ACTIVITY · {sp['minutes']} MIN", h=0.4, size=13, name="Activity tag")
    set_title(s, sp["title"], size=38, color=WHITE, x=0.8, y=1.5, w=11.5, h=1.0)
    n = len(sp["steps"])
    for i, st in enumerate(sp["steps"]):
        y = 2.85 + i * 0.85
        c = shape(s, MSO_SHAPE.OVAL, 0.9, y, 0.5, 0.5, PURPLE, name=f"Step {i + 1}")
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        r.font.name, r.font.size, r.font.bold = HEAD, Pt(17), True
        r.font.color.rgb = rgb(WHITE)
        text(s, 1.65, y - 0.05, 6.2, 0.8, dict(t=st, size=18, color=WHITE), name=f"Instruction {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.3, 2.85, 4.45, 3.7, PANEL, radius=0.06, name="Output card")
    text(s, 8.6, 3.05, 3.9, 0.4, dict(t="YOU LEAVE WITH", size=12, bold=True, color=CYAN), name="Output label")
    bullets(s, 8.6, 3.55, 3.9, 2.9, sp["output"], size=16, color=WHITE, gap=8, name="Output")
    return s


def checklist(prs, sp):
    s = light(prs, sp)
    items = sp["items"]
    n = len(items)
    per = math.ceil(n / 2)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.85, 12.15, 4.85, TINT, radius=0.04, shadow=True, name="Checklist card")
    rh = min(0.78, 4.4 / per)
    for i, it in enumerate(items):
        col, r = divmod(i, per)
        x = 0.95 + col * 6.1
        y = 2.1 + r * rh
        pic = s.shapes.add_picture(str(icon_file("check")), Inches(x), Inches(y + 0.08), Inches(0.36), Inches(0.36))
        pic.name = "icon:check"
        text(s, x + 0.55, y, 5.2, rh, dict(t=it, size=15), name=f"Item {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
    return s


def takeaway(prs, sp):
    s = dark(prs, sp)
    chip(s, 0.9, 1.3, 1.7, sp.get("tag", "REMEMBER"), h=0.4, size=13, name="Tag")
    text(s, 0.9, 2.2, 11.0, 2.8, dict(t=sp["text"], size=38, bold=True, font=HEAD, color=WHITE), name="Statement")
    if sp.get("sub"):
        text(s, 0.9, 5.2, 10.5, 1.3, dict(t=sp["sub"], size=20, color=SOFT), name="Sub")
    return s


def resources(prs, sp):
    s = light(prs, sp)
    groups = sp["groups"]
    n = len(groups)
    w = (12.15 - (n - 1) * 0.25) / n
    for i, g in enumerate(groups):
        x = 0.6 + i * (w + 0.25)
        d = i % 2 == 1
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.85, w, 4.85, NAVY if d else TINT, radius=0.06, shadow=True, name=f"Group {i + 1}")
        icon_in_circle(s, x + 0.3, 2.1, 0.65, g["icon"], NAVY_2 if d else WHITE)
        text(s, x + 0.3, 2.95, w - 0.6, 0.5, dict(t=g["head"], size=19, bold=True, font=HEAD, color=WHITE if d else INK), name=f"Group {i + 1} head")
        text(s, x + 0.3, 3.55, w - 0.6, 3.0, [dict(t=l, size=14, color=SOFT if d else SLATE, space_after=7) for l in g["lines"]], name=f"Group {i + 1} lines")
    return s


def stat(prs, sp):
    s = dark(prs, sp)
    text(s, 0.9, 1.5, 6.0, 2.6, dict(t=sp["stat"], size=130, bold=True, font=HEAD, color=AMBER), name="Stat")
    text(s, 0.9, 4.1, 5.8, 0.9, dict(t=sp["label"], size=22, bold=True, color=WHITE), name="Label")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.3, 1.5, 5.4, 4.9, PANEL, radius=0.05, name="Points card")
    bullets(s, 7.65, 1.85, 4.8, 4.3, sp["points"], size=17, color=WHITE, gap=12, name="Points")
    if sp.get("source"):
        text(s, 0.9, 6.7, 11.5, 0.3, dict(t="Source: " + sp["source"], size=11, color=MUTED), name="Source")
    return s


TEMPLATES = {f.__name__: f for f in (cover, section, cards, compare, steps, vsteps, bullets_stat, two_col, code, table, activity, checklist, takeaway, resources, stat)}


def build_deck(specs: list[dict], out: Path) -> Path:
    global _NUMBER_SP
    prs = Presentation(SRC)
    donor = next(sl for sl in prs.slides if sl.shapes.title is not None and sl.shapes.title.text_frame.text.strip() == "What we'll cover today")
    for sh in donor.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == 13:
            _NUMBER_SP = copy.deepcopy(sh._element)
    for sl in list(prs.slides):
        delete_slide(prs, sl)
    prs.part.rename_slide_parts([sld.rId for sld in prs.slides._sldIdLst])
    for i, sp in enumerate(specs, 1):
        TEMPLATES[sp["kind"]](prs, sp)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
