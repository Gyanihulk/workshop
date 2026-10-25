"""Turn the slides built by build.py into Google Slides API batchUpdate requests.

    python3 to_slides.py            # writes out/slides_api/*.json

The slides are not drawn a second time: this reads out/workshop-deck.pptx, so Google gets
exactly what was rendered and checked locally. Pictures are referenced by name and resolved
through out/slides_api/urls.json (name -> image URL taken from icons already in the deck).

Order of operations (see README): 1) structure.json duplicates donor slides so the layout,
title placeholder and slide-number placeholder come along, 2) one <slide>.json per slide.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

HERE = Path(__file__).parent
PPTX = HERE / "out" / "workshop-deck.pptx"
OUT = HERE / "out" / "slides_api"

# slide title in the pptx -> (new slide id, donor slide in the Google deck, donor child ids to keep)
SLIDES = {
    "Hi, I'm Adamya": dict(id="n_about", donor="p1", keep_title="p1_i5", keep_number=None),
    "Four jobs, four lessons": dict(id="n_exp", donor="p3", keep_title="p3_i2", keep_number="p3_i25"),
    "Beyond the job": dict(id="n_beyond", donor="p3", keep_title="p3_i2", keep_number="p3_i25"),
    "Warm-up: pick your corner": dict(id="n_warm", donor="p3", keep_title="p3_i2", keep_number="p3_i25"),
    "From classroom to first client": dict(id="n_flow", donor="p3", keep_title="p3_i2", keep_number="p3_i25"),
}
DONOR_CHILDREN = {
    "p1": [f"p1_i{k}" for k in range(2, 11)],
    "p3": [f"p3_i{k}" for k in list(range(2, 23)) + [25]],
}
PRST = {"rect": "RECTANGLE", "roundRect": "ROUND_RECTANGLE", "ellipse": "ELLIPSE", "triangle": "TRIANGLE"}
INSET_X, INSET_Y = 91440, 45720  # default text-box insets in Google Slides (0.1in, 0.05in)


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return {k: round(int(h[i:i + 2], 16) / 255, 3) for k, i in (("red", 0), ("green", 2), ("blue", 4))}


def xform(x, y, w, h, rot=0.0):
    """Absolute transform for an element whose base size is (w,h) EMU, rotated rot degrees about its centre."""
    if not rot:
        return dict(scaleX=1, scaleY=1, translateX=x, translateY=y, unit="EMU")
    t = math.radians(rot)
    c, s = math.cos(t), math.sin(t)
    cx, cy = x + w / 2, y + h / 2
    return dict(
        scaleX=round(c, 6), scaleY=round(c, 6), shearX=round(-s, 6), shearY=round(s, 6),
        translateX=round(cx - (c * w / 2 - s * h / 2)), translateY=round(cy - (s * w / 2 + c * h / 2)), unit="EMU",
    )


def props(w, h, x, y, rot=0.0, page=None):
    return {
        "pageObjectId": page,
        "size": {"width": {"magnitude": int(w), "unit": "EMU"}, "height": {"magnitude": int(h), "unit": "EMU"}},
        "transform": xform(int(x), int(y), int(w), int(h), rot),
    }


def fill_hex(sh):
    try:
        return str(sh.fill.fore_color.rgb)
    except Exception:
        return None


def shadow_of(sh):
    eff = sh._element.spPr.find(qn("a:effectLst"))
    return eff is not None and eff.find(qn("a:outerShdw")) is not None


def text_requests(oid, tf):
    """insertText + styles for every run, alignment and vertical anchor."""
    paras = tf.paragraphs
    full, runs, aligns, idx = [], [], [], 0
    for pi, p in enumerate(paras):
        start = idx
        for r in p.runs:
            full.append(r.text)
            runs.append((idx, idx + len(r.text), r))
            idx += len(r.text)
        pPr = p._p.pPr
        has_bullet = pPr is not None and pPr.find(qn("a:buChar")) is not None
        aligns.append((start, idx, p.alignment, has_bullet, p.space_after))
        if pi < len(paras) - 1:
            full.append("\n")
            idx += 1
    text = "".join(full)
    if not text:
        return []
    out = [{"insertText": {"objectId": oid, "text": text}}]
    for s, e, r in runs:
        f = r.font
        style, fields = {}, []
        if f.name:
            style["fontFamily"] = f.name
            fields.append("fontFamily")
        if f.size:
            style["fontSize"] = {"magnitude": round(f.size.pt, 1), "unit": "PT"}
            fields.append("fontSize")
        if f.bold:
            style["bold"] = True
            fields.append("bold")
        if f.italic:
            style["italic"] = True
            fields.append("italic")
        try:
            style["foregroundColor"] = {"opaqueColor": {"rgbColor": rgb(str(f.color.rgb))}}
            fields.append("foregroundColor")
        except Exception:
            pass
        out.append({"updateTextStyle": {"objectId": oid, "textRange": {"type": "FIXED_RANGE", "startIndex": s, "endIndex": e}, "style": style, "fields": ",".join(fields)}})
    for s, e, a, bullet, space in aligns:
        if space is not None:
            out.append({"updateParagraphStyle": {"objectId": oid, "textRange": {"type": "FIXED_RANGE", "startIndex": s, "endIndex": max(e, s + 1)}, "style": {"spaceBelow": {"magnitude": round(space.pt, 1), "unit": "PT"}}, "fields": "spaceBelow"}})
        if bullet:
            out.append({"createParagraphBullets": {"objectId": oid, "textRange": {"type": "FIXED_RANGE", "startIndex": s, "endIndex": max(e, s + 1)}, "bulletPreset": "BULLET_DISC_CIRCLE_SQUARE"}})
        al = {PP_ALIGN.CENTER: "CENTER", PP_ALIGN.RIGHT: "END"}.get(a)
        if al:
            out.append({"updateParagraphStyle": {"objectId": oid, "textRange": {"type": "FIXED_RANGE", "startIndex": s, "endIndex": max(e, s + 1)}, "style": {"alignment": al}, "fields": "alignment"}})
    anchor = {MSO_ANCHOR.MIDDLE: "MIDDLE", MSO_ANCHOR.BOTTOM: "BOTTOM"}.get(tf.vertical_anchor)
    if anchor:
        out.append({"updateShapeProperties": {"objectId": oid, "shapeProperties": {"contentAlignment": anchor}, "fields": "contentAlignment"}})
    return out


def shadow_props():
    return {
        "shadow": {
            "type": "OUTER", "alignment": "BOTTOM_LEFT", "alpha": 0.14, "rotateWithShape": False,
            "blurRadius": {"magnitude": 84666, "unit": "EMU"},
            "color": {"rgbColor": {"red": 0.1176, "green": 0.1529, "blue": 0.3804}},
            "transform": {"scaleX": 1, "scaleY": 1, "translateX": 0, "translateY": 25400, "unit": "EMU"},
        }
    }


def convert(slide, cfg):
    sid = cfg["id"]
    reqs, n = [], 0

    def nid():
        nonlocal n
        n += 1
        return f"{sid}_{n:02d}"

    # remove everything the donor slide brought, except the title (and slide-number) placeholder
    keep = {cfg["keep_title"], cfg["keep_number"]} - {None}
    for child in DONOR_CHILDREN[cfg["donor"]]:
        if child not in keep:
            reqs.append({"deleteObject": {"objectId": f"{sid}_{child.split('_', 1)[1]}"}})
    title_id = f"{sid}_{cfg['keep_title'].split('_', 1)[1]}"

    for sh in slide.shapes:
        if sh.is_placeholder:
            if sh.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                reqs.append({"deleteText": {"objectId": title_id, "textRange": {"type": "ALL"}}})
                reqs += text_requests(title_id, sh.text_frame)
                if True:
                    reqs.append({"updatePageElementTransform": {"objectId": title_id, "applyMode": "ABSOLUTE", "transform": dict(scaleX=round(sh.width / 3000000, 5), scaleY=round(sh.height / 3000000, 5), translateX=int(sh.left), translateY=int(sh.top), unit="EMU")}})
            continue
        x, y, w, h = int(sh.left), int(sh.top), int(sh.width), int(sh.height)
        oid = nid()
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            reqs.append({"createImage": {"objectId": oid, "url": "ICON:" + sh.name.removeprefix("icon:"), "elementProperties": props(w, h, x, y, page=sid)}})
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            reqs.append({"createShape": {"objectId": oid, "shapeType": "TEXT_BOX", "elementProperties": props(w + 2 * INSET_X, h + 2 * INSET_Y, x - INSET_X, y - INSET_Y, page=sid)}})
            reqs += text_requests(oid, sh.text_frame)
            continue
        prst = sh._element.spPr.find(qn("a:prstGeom")).get("prst")
        reqs.append({"createShape": {"objectId": oid, "shapeType": PRST[prst], "elementProperties": props(w, h, x, y, sh.rotation, page=sid)}})
        reqs.append({"updateShapeProperties": {"objectId": oid, "fields": "shapeBackgroundFill.solidFill.color,outline.propertyState", "shapeProperties": {"shapeBackgroundFill": {"solidFill": {"color": {"rgbColor": rgb(fill_hex(sh))}}}, "outline": {"propertyState": "NOT_RENDERED"}}}})
        if shadow_of(sh):
            reqs.append({"updateShapeProperties": {"objectId": oid, "fields": "shadow", "shapeProperties": shadow_props()}})
        if sh.has_text_frame and sh.text_frame.text:
            reqs += text_requests(oid, sh.text_frame)
    return reqs


def main():
    import sys

    only = set(sys.argv[1:])  # e.g. python3 to_slides.py n_about n_exp n_beyond
    prs = Presentation(PPTX)
    OUT.mkdir(parents=True, exist_ok=True)
    chosen = {title: cfg for title, cfg in SLIDES.items() if not only or cfg["id"] in only}
    structure = []
    # duplicates land right after their donor: later duplicates of p3 end up first, so list them in reverse order
    for title, cfg in reversed(list(chosen.items())):
        children = {c: f"{cfg['id']}_{c.split('_', 1)[1]}" for c in DONOR_CHILDREN[cfg["donor"]]}
        structure.append({"duplicateObject": {"objectId": cfg["donor"], "objectIds": {cfg["donor"]: cfg["id"], **children}}})
    (OUT / "structure.json").write_text(json.dumps(structure, separators=(",", ":")))
    for sl in prs.slides:
        t = sl.shapes.title.text_frame.text if sl.shapes.title is not None else ""
        if t in chosen:
            cfg = chosen[t]
            data = convert(sl, cfg)
            (OUT / f"{cfg['id']}.json").write_text(json.dumps(data, separators=(",", ":")))
            print(cfg["id"], len(data), "requests")
    print("structure:", len(structure), "requests")


if __name__ == "__main__":
    main()
